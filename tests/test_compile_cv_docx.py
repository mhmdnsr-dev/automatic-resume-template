from __future__ import annotations

import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document

from helpers.compile_cv_docx import _inline_segments, compile_docx, parse_resume


ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "resume.tex"


class DocxCompilerTests(unittest.TestCase):
    def test_inline_segments_preserve_link_position_and_en_dash(self) -> None:
        segments = _inline_segments(
            r"Contributed to \textbf{\href{https://example-booking.com}{example-booking.com}}, "
            r"processing 100K--200K searches."
        )
        text = "".join(segment[1] for segment in segments)
        self.assertEqual(
            text,
            "Contributed to example-booking.com, processing 100K\u2013200K searches.",
        )
        self.assertEqual(
            [segment for segment in segments if segment[0] == "link"],
            [("link", "example-booking.com", "https://example-booking.com")],
        )

    def test_parses_master_sections_from_tex(self) -> None:
        resume = parse_resume(MASTER)
        self.assertEqual(resume.name, "Example Candidate")
        self.assertIn("Frontend Developer", resume.summary)
        self.assertGreaterEqual(len(resume.experience), 2)
        self.assertGreaterEqual(len(resume.projects), 2)
        self.assertGreaterEqual(len(resume.education), 2)
        self.assertGreaterEqual(len(resume.skills), 5)

    def test_compiles_valid_docx_with_links_and_real_bullets(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / MASTER.name
            source.write_text(MASTER.read_text(encoding="utf-8"), encoding="utf-8")
            output = compile_docx(source)
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("EXAMPLE CANDIDATE", text)
            self.assertIn("Professional Summary", text)
            self.assertIn("Technical Skills", text)
            self.assertIn("100K\u2013200K", text)
            self.assertGreaterEqual(
                sum(paragraph.style.name == "List Bullet" for paragraph in document.paragraphs),
                8,
            )
            with ZipFile(output) as package:
                relationships = package.read("word/_rels/document.xml.rels").decode("utf-8")
                numbering = package.read("word/numbering.xml").decode("utf-8")
                app_properties = ElementTree.fromstring(package.read("docProps/app.xml"))
            self.assertEqual(relationships.count('TargetMode="External"'), 13)
            self.assertNotIn("bit.ly", relationships)
            self.assertIn("https://github.com/example-org/event-booking-platform/", relationships)
            self.assertIn("https://example-events.app/", relationships)
            self.assertIn("<w:abstractNum ", numbering)

            name = next(paragraph for paragraph in document.paragraphs if paragraph.text == "EXAMPLE CANDIDATE")
            self.assertEqual(
                document.paragraphs[1].text,
                "Remote City | +10000000000 | candidate@example.com",
            )
            self.assertEqual(
                document.paragraphs[2].text,
                "linkedin.com/in/example-candidate | github.com/example-candidate | candidate.dev",
            )
            summary = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("Frontend Developer with 2.5+ years")
            )
            section = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Experience")
            bullet = next(paragraph for paragraph in document.paragraphs if paragraph.style.name == "List Bullet")
            company = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("Example Travel Platform")
            )
            skills_row = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("Languages:")
            )
            self.assertEqual(name.runs[0].font.size.pt, 24.0)
            self.assertEqual(section.runs[0].font.size.pt, 12.0)
            self.assertEqual(summary.runs[0].font.size.pt, 10.0)
            self.assertEqual(bullet.runs[0].font.size.pt, 10.0)
            self.assertEqual(company.runs[0].font.size.pt, 11.0)
            self.assertEqual(skills_row.runs[0].font.size.pt, 10.0)
            self.assertEqual(skills_row.paragraph_format.line_spacing.pt, 10.5)

            employer_subtitle = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith(
                    "Frontend Developer\tHybrid (Remote City / Remote)"
                )
            )
            education_subtitle = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("Intensive Code Camp")
            )
            self.assertTrue(employer_subtitle._p.xpath("./w:pPr/w:keepNext"))
            education_keep_next = education_subtitle._p.xpath("./w:pPr/w:keepNext")
            self.assertEqual(len(education_keep_next), 1)
            self.assertEqual(education_keep_next[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"), "0")

            standalone_links = {
                "example-booking.com",
                "example-games.com",
                "example-shop.com",
                "example-writer.app",
                "GitHub",
                "Live",
            }
            self.assertFalse(
                any(paragraph.text in standalone_links for paragraph in document.paragraphs)
            )
            linked_bullets = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.style.name == "List Bullet"
                and paragraph._p.xpath(".//w:hyperlink")
            ]
            self.assertEqual(len(linked_bullets), 6)
            self.assertTrue(
                any(paragraph.text.startswith("Contributed to example-booking.com,") for paragraph in linked_bullets)
            )
            self.assertTrue(
                any("Links: GitHub / Live" in paragraph.text for paragraph in linked_bullets)
            )

            section_properties = document.sections[0]
            self.assertAlmostEqual(section_properties.top_margin.inches, 0.58, places=2)
            self.assertAlmostEqual(section_properties.bottom_margin.inches, 0.58, places=2)
            self.assertAlmostEqual(section_properties.left_margin.inches, 0.58, places=2)
            self.assertAlmostEqual(section_properties.right_margin.inches, 0.58, places=2)

            props = document.core_properties
            self.assertEqual(props.title, "Example Candidate - Frontend Developer")
            self.assertEqual(props.author, "Example Candidate")
            self.assertEqual(
                props.subject,
                "Frontend Developer Resume - React and Angular",
            )
            self.assertEqual(
                props.comments,
                "Example Candidate frontend resume generated from the canonical LaTeX source",
            )
            self.assertEqual(props.last_modified_by, "Example Candidate")
            self.assertEqual(props.created.date(), datetime.now(timezone.utc).date())
            self.assertEqual(props.modified.date(), datetime.now(timezone.utc).date())

            namespace = {
                "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
            }
            self.assertEqual(app_properties.findtext("ep:Pages", namespaces=namespace), "1")
            self.assertGreater(
                int(app_properties.findtext("ep:Words", default="0", namespaces=namespace)),
                0,
            )
            self.assertGreater(
                int(app_properties.findtext("ep:Characters", default="0", namespaces=namespace)),
                0,
            )


if __name__ == "__main__":
    unittest.main()
